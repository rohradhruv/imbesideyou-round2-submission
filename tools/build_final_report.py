"""Build a polished DOCX rendering of FINAL_REPORT.md for submission."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "FINAL_REPORT.md"
DESTINATION = ROOT / "FINAL_REPORT.docx"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def font(run, size=10, bold=False, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    font(run, 8.5, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def parse_row(line):
    return [part.strip() for part in line.strip().strip("|").split("|")]


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(line[2:]), 20, True, (0, 0, 0))
        elif line.startswith("## "):
            p = doc.add_heading(line[3:], level=1)
            for run in p.runs:
                font(run, 14, True, (0, 0, 0))
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], level=2)
            for run in p.runs:
                font(run, 11, True, (0, 0, 0))
        elif line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|---"):
            headers = parse_row(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, headers):
                shade(cell, "1F4E78")
                add_text(cell, header, True, (255, 255, 255))
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    if row_index % 2:
                        shade(cell, "F2F6FA")
                    add_text(cell, value.replace("`", ""))
            continue
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            font(p.add_run(line.replace("`", "").replace("**", "")), 10)
        i += 1
    doc.save(DESTINATION)
    print(DESTINATION)


if __name__ == "__main__":
    main()
