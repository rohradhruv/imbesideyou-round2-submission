"""Review-first prototype for Dataset B HR new-hire onboarding verification.

The automation prepares deterministic checks and a reviewer-ready checklist.
It never completes or flags a production HR case: those decisions remain human
approval steps because the logs do not establish authoritative policy or access.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = REPO_ROOT / "prototype_output"
AUDIT_LOG = OUTPUT_DIR / "onboarding_automation_log.csv"


@dataclass(frozen=True)
class OnboardingItem:
    code: str
    label: str
    required: bool
    evidence_received: bool
    source_status: str
    note: str = ""


@dataclass(frozen=True)
class OnboardingCase:
    case_id: str
    employee_id: str
    employee_name: str
    department: str
    start_date: str
    hire_type: str
    items: tuple[OnboardingItem, ...]


# Local mock adapter: production implementation would call the HR system under
# approved service credentials and map its authoritative item/status fields.
MOCK_CASES = {
    "ONB-1001": OnboardingCase(
        "ONB-1001", "E2288", "Sakura Uchida", "Administration", "2026-07-26", "new_graduate",
        (
            OnboardingItem("MY_NUMBER", "My Number submission and verification", True, True, "verified"),
            OnboardingItem("HEALTH", "Health insurance and pension enrollment", True, True, "verified"),
            OnboardingItem("BANK", "Payroll bank account notification", True, True, "verified"),
            OnboardingItem("COMMUTE", "Commuting allowance application", True, True, "verified"),
            OnboardingItem("EMPLOYMENT_INSURANCE", "Employment insurance certificate", False, False, "not_applicable", "Not required for new graduate hire."),
        ),
    ),
    "ONB-1002": OnboardingCase(
        "ONB-1002", "E2294", "Haruto Suzuki", "Operations", "2026-07-11", "experienced_hire",
        (
            OnboardingItem("MY_NUMBER", "My Number submission and verification", True, True, "verified"),
            OnboardingItem("HEALTH", "Health insurance and pension enrollment", True, True, "verified"),
            OnboardingItem("BANK", "Payroll bank account notification", True, False, "pending"),
            OnboardingItem("COMMUTE", "Commuting allowance application", True, True, "verified"),
            OnboardingItem("EMPLOYMENT_INSURANCE", "Employment insurance certificate", True, False, "missing"),
        ),
    ),
    "ONB-1003": OnboardingCase(
        "ONB-1003", "E2302", "Kenji Mori", "Systems", "2026-07-11", "experienced_hire",
        (
            OnboardingItem("MY_NUMBER", "My Number submission and verification", True, True, "verified"),
            OnboardingItem("HEALTH", "Health insurance and pension enrollment", True, True, "verified"),
            OnboardingItem("BANK", "Payroll bank account notification", True, True, "verified"),
            OnboardingItem("COMMUTE", "Commuting allowance application", True, True, "verified"),
            OnboardingItem("EMPLOYMENT_INSURANCE", "Employment insurance certificate", True, True, "conflicting", "Portal status conflicts with attached evidence; verify manually."),
        ),
    ),
}


def evaluate_item(item: OnboardingItem) -> tuple[str, str]:
    """Apply only transparent deterministic rules; do not infer HR policy."""
    if not item.required:
        return "NOT APPLICABLE", item.note or "Not required for this hire type."
    if item.source_status == "verified" and item.evidence_received:
        return "PASS", "Required evidence is recorded as verified."
    if item.source_status == "conflicting":
        return "REQUIRES REVIEW", item.note or "Source data is inconsistent."
    if not item.evidence_received or item.source_status in {"missing", "pending"}:
        return "MISSING", item.note or "Required evidence is absent or not yet verified."
    return "REQUIRES REVIEW", item.note or "Unexpected source state."


def evaluate_case(case: OnboardingCase) -> tuple[list[tuple[OnboardingItem, str, str]], str, str]:
    results = [(item, *evaluate_item(item)) for item in case.items]
    exceptions = [result for result in results if result[1] in {"MISSING", "REQUIRES REVIEW"}]
    if not exceptions:
        return results, "READY FOR HUMAN COMPLETION", "All required onboarding items passed deterministic checks."
    if any(result[1] == "MISSING" for result in exceptions):
        return results, "HUMAN REVIEW REQUIRED", "Missing required evidence must be resolved before completion."
    return results, "HUMAN REVIEW REQUIRED", "A source-data conflict requires reviewer judgment before completion."


def set_cell_fill(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def set_run_font(run, size=10, bold=False, color=None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_text(cell, text: str, bold=False, color=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_run_font(run, 9, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_heading(doc, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run, 14 if level == 1 else 11, bold=True, color=(0, 0, 0))


def build_checklist(case: OnboardingCase, results, decision: str, rationale: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(title.add_run("New Hire Onboarding Review Checklist"), 18, bold=True, color=(0, 0, 0))
    intro = doc.add_paragraph()
    set_run_font(intro.add_run("Automation-prepared review. A human reviewer must decide whether to complete or flag the HR case."), 10)

    add_heading(doc, "Case summary", 1)
    summary = doc.add_table(rows=0, cols=2)
    summary.style = "Table Grid"
    for label, value in [
        ("Case ID", case.case_id), ("Employee", f"{case.employee_name} ({case.employee_id})"),
        ("Department", case.department), ("Start date", case.start_date), ("Hire type", case.hire_type.replace("_", " ")),
        ("Automation recommendation", decision),
    ]:
        cells = summary.add_row().cells
        set_cell_fill(cells[0], "D9E2F3")
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], value, bold=(label == "Automation recommendation"))

    add_heading(doc, "Required item checks", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Item", "Requirement", "Automation result", "Evidence and next action"]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_fill(cell, "1F4E78")
        set_cell_text(cell, header, bold=True, color=(255, 255, 255))

    colors = {"PASS": "E2F0D9", "MISSING": "FCE4D6", "REQUIRES REVIEW": "FFF2CC", "NOT APPLICABLE": "EDEDED"}
    for item, result, detail in results:
        cells = table.add_row().cells
        set_cell_text(cells[0], item.label)
        set_cell_text(cells[1], "Required" if item.required else "Not required")
        set_cell_fill(cells[2], colors[result])
        set_cell_text(cells[2], result, bold=True)
        set_cell_text(cells[3], detail)

    add_heading(doc, "Suggested verification note", 1)
    note = doc.add_paragraph()
    if decision == "READY FOR HUMAN COMPLETION":
        text = f"Onboarding case {case.case_id}: required items were checked and no deterministic exceptions were found. Please verify the source records and complete the case if approved."
    else:
        text = f"Onboarding case {case.case_id}: {rationale} Review the exceptions above, then either resolve the case or flag it for follow-up."
    set_run_font(note.add_run(text), 10)

    add_heading(doc, "Reviewer action", 1)
    action = doc.add_paragraph()
    set_run_font(action.add_run("Reviewer decision:  [ ] Complete case    [ ] Flag for follow-up    [ ] Hold pending evidence"), 10, bold=True)
    signoff = doc.add_paragraph()
    set_run_font(signoff.add_run("Reviewed by: ____________________    Date: ____________________"), 10)

    safe_id = re.sub(r"[^A-Za-z0-9_-]+", "_", case.case_id)
    path = OUTPUT_DIR / f"onboarding_review_{safe_id}.docx"
    doc.save(path)
    return path


def log_run(case_id: str, decision: str, output: Path, status: str, note: str) -> None:
    new_file = not AUDIT_LOG.exists()
    with AUDIT_LOG.open("a", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        if new_file:
            writer.writerow(["timestamp", "case_id", "automation_recommendation", "output_file", "status", "note"])
        writer.writerow([datetime.now().isoformat(timespec="seconds"), case_id, decision, output.name, status, note])


def process_case(case_id: str) -> bool:
    case = MOCK_CASES.get(case_id)
    if case is None:
        log_run(case_id, "N/A", Path("-"), "failed", "Unknown onboarding case ID.")
        print(f"ERROR {case_id}: unknown onboarding case ID")
        return False

    results, decision, rationale = evaluate_case(case)
    output = build_checklist(case, results, decision, rationale)
    log_run(case.case_id, decision, output, "generated_pending_human_review", rationale)
    print(f"{case.case_id}: {decision} -> {output.name}")
    return True


def main() -> int:
    parser = argparse.ArgumentParser(description="HR new-hire onboarding review assistant")
    parser.add_argument("case_ids", nargs="*", help="Case IDs; defaults to all local mock cases")
    args = parser.parse_args()
    case_ids = args.case_ids or list(MOCK_CASES)
    print("HR new-hire onboarding review assistant")
    print("Human portal completion or flagging is always required.")
    succeeded = sum(process_case(case_id) for case_id in case_ids)
    print(f"Completed: {succeeded}/{len(case_ids)}")
    return 0 if succeeded == len(case_ids) else 1


if __name__ == "__main__":
    sys.exit(main())
